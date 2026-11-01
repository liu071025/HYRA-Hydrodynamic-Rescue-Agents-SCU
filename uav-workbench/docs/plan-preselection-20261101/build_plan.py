from pathlib import Path
import json, hashlib
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

BASE=Path(__file__).parent
DOC=Document()
s=DOC.sections[0];s.page_width=Inches(8.5);s.page_height=Inches(11)
s.top_margin=s.bottom_margin=Inches(.7);s.left_margin=s.right_margin=Inches(.75)
for name in ('Normal','Title','Subtitle','Heading 1','Heading 2'):
 st=DOC.styles[name];st.font.name='Microsoft YaHei';st.font.color.rgb=RGBColor(0,0,0)
 st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
 st.font.size=Pt(11 if name=='Normal' else 23 if name=='Title' else 16 if name=='Heading 1' else 12)
 st.paragraph_format.space_after=Pt(7)
 st.paragraph_format.line_spacing=1.17
DOC.styles['Normal'].paragraph_format.widow_control=True
header=s.header.paragraphs[0];header.text='HYRA  无人机集群搜救研究';header.style='Caption'
footer=s.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('研究计划 v1.0  |  2026年11月1日  |  ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
content=[]
def p(text):DOC.add_paragraph(text);content.append(text)
def h(text):DOC.add_heading(text,1);content.append(text)
def page(title):DOC.add_page_break();h(title)
def table(headers,rows,widths):
 t=DOC.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for col,w in zip(t.columns,widths):col.width=Inches(w)
 for i,x in enumerate(headers):t.rows[0].cells[i].text=x
 for row in rows:
  cells=t.add_row().cells
  for i,x in enumerate(row):cells[i].text=str(x)
 for ri,row in enumerate(t.rows):
  row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
  if ri==0:row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
  for ci,cell in enumerate(row.cells):
   cell.width=Inches(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   pr=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
   for side in ('top','left','bottom','right'):
    b=OxmlElement('w:'+side);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
   pr.append(borders);m=OxmlElement('w:tcMar')
   for side in ('top','left','bottom','right'):
    v=OxmlElement('w:'+side);v.set(qn('w:w'),'100');v.set(qn('w:type'),'dxa');m.append(v)
   pr.append(m);sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'DCE6F1' if ri==0 else 'FFFFFF' if ri%2 else 'F5F7F9');pr.append(sh)
   for para in cell.paragraphs:
    para.paragraph_format.space_after=Pt(2);para.paragraph_format.line_spacing=1.12
    for r in para.runs:r.font.size=Pt(10.5);r.font.bold=(ri==0);r.font.color.rgb=RGBColor(0,0,0)
 content.append({'headers':headers,'rows':rows})
 DOC.add_paragraph().paragraph_format.space_after=Pt(1)

DOC.add_paragraph('无人机搜救算法预选与训练计划书','Title')
p('版本 v1.0    冻结日期 2026年11月1日    对象 HYRA 项目研究与复核')
p('本计划先完成失败回放和五候选预选，再按证据决定是否训练新的分层策略。推荐研究主线为 HATRPO 任务分配加覆盖规划，双学习器候选为 HATRPO 任务分配加 HASAC 局部导航。首轮执行仅更新诊断和选择记录，使用冻结模型；后续学习设有明确进入条件。')
h('1 当前依据与研究目标')
table(['已有结果','观察','对决策的影响'],[
 ['较强单模型','HATRPO1024 加共同施救和重试控制，覆盖96.65%，获救99.38%，整场完成0/16','保留为强对照，不能只和较弱组合比'],
 ['联合训练','HASAC加HATRPO保留64回合版：获救94.53%，覆盖94.45%；256回合冲突及较弱分工退化','已有组合未证明优于单模型'],
 ['覆盖补漏','六种256回合模型加补漏在旧8次验证均完成；部分候选冲突仍高','先查规划层，明确接管比例与规则贡献']],[1.22,3.05,2.73])
p('主目标是在固定300个联合步内完成全部实际观测和全部20名目标的模拟服务，并降低完成时间及潜在冲突。水情正确性、模拟服务完成和真实救援能力是不同结论。当前数据为合成水场，本计划不宣称现场搜救有效。')
p('禁止使用 MAPPO、QMIX、VDN、IPPO、DQN、DDPG。旧模型、坏结果、原环境及封存50场测试保留。最多一项重负载，CPU两线程串行；不自动推送、发消息或改变系统安全设置。')

page('2 候选路线与预选对象')
table(['路线','职责划分','本轮定位'],[
 ['纯规划基线','已观测历史重建加分区搜索、补漏、发现后分配和连续服务','必做；量化不使用神经模型时的表现'],
 ['HATRPO 加覆盖规划','新高层选择搜索区或已发现目标，规划器实现低层移动和协调','优先研究方向；高层接口尚未实现'],
 ['HATRPO 加 HASAC','高层分工，低层按公开目标和历史观测导航，统一服务和安全规则','备选；需先证明层级接口和训练有效']],[1.5,3.6,1.9])
p('当前 HATRPO 和 HASAC 权重直接输出逐步动作，不能直接冒充高层任务分配器。本轮先筛选可复用的低层基线及规划组合；任何新高层都须另登记输入、动作、奖励与检查点身份。')
table(['预选编号','冻结策略','覆盖规划'],[
 ['P0','纯规划加相同服务和重试控制，零模型训练','从第1步启用'],
 ['P1','HATRPO1024 加服务和 RetryGuard','关闭，作强单模型对照'],
 ['P2','同一 HATRPO1024 权重及控制','连续5步无新增覆盖后启用'],
 ['P3','HASAC256 加同一服务和 RetryGuard','相同补漏条件'],
 ['P4','MAAC256 加同一服务和 RetryGuard','相同补漏条件']],[.65,4.25,2.1])
p('选择 MAAC 为补充候选，是因为其旧小批补漏组合冲突较低，并非证明 MAAC 最优。暂不扩大 MAT 或 COMA 的长训练；其历史失败保留。联合64回合候选继续归档，本轮不以增加组合数量代替机制诊断。')

page('3 共同环境与信息权限')
table(['项目','冻结口径'],[
 ['空间和资源','1000米见方、100×100格、4架无人机、20名固定目标；300步，每步5秒'],
 ['感知和运动','180度扇形、半径100米、仅决策时刻采样；9种离散动作，正常每步50米路径预算'],
 ['模拟服务','已发现目标10米范围内，同机连续悬停10秒；水深不高于0.08米、流速不高于0.3米每秒'],
 ['通信假设','理想全连接共享已允许的观测和自身状态；没有验证丢包、带宽、建筑或真实续航'],
 ['完成条件','全部10000格实际被观察且20人全部完成模拟服务；不得降低覆盖要求或延长300步'],
 ['冲突计数','原始意图的轨迹潜在冲突，按机对计数；环境拦停后的零碰撞不抹除原始冲突']],[1.25,5.75])
p('决策仅使用当时已取得的观测、公开位置朝向、已发现目标和观察历史。未知水情保持未知；不可读隐藏目标、未来水情或后台完整覆盖真值。洪水受淹不自动构成无人机禁飞区。')
p('评价器可以在回合结束后读取真值，核对遗漏位置、服务窗口和失败类型，但这些诊断结果不能回流到同一回合控制器。复现旧对照时保持服务控制、重试控制、动作选择和随机种子一致。')
p('首轮逐步保存原策略动作、规划后动作、服务状态、位移、新增覆盖、目标发现与服务事件、潜在冲突及拦停。每回合保留全部20人最终状态、漏扫格坐标和未完成原因；不删除失败目标或失败场景。')

page('4 首轮回放与预选协议')
p('复用2个既有合成验证父水场。旧布局固定为9001至9004，新增布局固定为9401至9404。新增布局只能说明同一水场中的布局适应性，不能称新的洪水事件或独立最终测试。冻结后再运行候选，原场景及源码和模型哈希随协议保存。')
table(['执行项','次数与口径'],[
 ['P0','2父水场×8布局，固定规划，16次'],
 ['P1 至 P4 随机执行','每候选2父水场×8布局×2动作种子3101和3102，32次'],
 ['P1 至 P4 固定执行','每候选2父水场×8布局，每次取概率最高动作，16次'],
 ['合计','208次回放，新增训练0回合；动作种子不是训练种子']],[2.1,4.9])
p('先作独立短检查：观测重建的覆盖不能超过环境实际覆盖，强制服务动作必须满足已登记掩码，评价记录不得改变策略动作，源码和父权重不能被更新；旧布局P1随机回放必须复现已存对应指标。任何不一致先定位，失败日志保留，不以改结果文件解决。')
p('失败类型分别记录：未发现；已发现未完成服务且从发现时刻起存在连续合法服务窗口；已发现未完成且没有这样的窗口。后一类仅表示发现后的窗口缺失，不证明整个任务不可行。另记冲突拦停、无新增覆盖、普通悬停及施救中断。可重叠原因不能相加冒充独立失败人数。')
p('旧布局、新布局分别报告。随机方式先在每个父水场和布局上平均动作种子，再汇总；固定方式单列。指标包含整场完成、获救、覆盖、300步截断完成时间、首次服务、全员服务时间、潜在冲突与规划接管比例。')

page('5 预选条件与停止规则')
p('本轮首要决定是是否值得继续扩大当前逐步策略组合。分层研究方向的预选仅为待验证假设，不因冻结权重加补漏有效就宣称高层学习有效。以下条件在预选结果出现前固定。')
table(['门槛','预定要求'],[
 ['数值与信息','参数有限、动作合法、原数据和源码哈希一致、无控制器读取隐藏真值；全部必要检查通过'],
 ['工程可用候选','在旧与新增布局各自的随机和固定执行中，完成比例和获救比例均不低于P0，平均潜在冲突每步不高于0.02'],
 ['学习值得投入','满足工程门槛，并在旧与新两组随机执行中分别增加至少一个场景等效的完成量，或同等完成和获救下平均截断步数少至少5步'],
 ['并列选择','整场完成、获救、截断步数、冲突依次比较，仍相同选计算量更低者；结果全保留'],
 ['无候选通过','不追加原方案1024回合；保留纯规划及强单模型，转为记录具体失效机制和下一项单因素接口试验']],[1.55,5.45])
p('0.02是本研究开发筛选阈值，不是实物飞行安全标准。平均冲突之外保留每场原始计数及最大值；高均值获救不能掩盖差场景或冲突退步。失败回合的截断时间记300步，禁止只在成功回合上挑选更短时间。')
p('每项修复先写明改动、物理或优化理由、缺口、最小试验、成本和否决条件。奖励、学习率、控制接管和场景扩展分开改，避免同时改动后无法判断原因。预选无收益不等于分层方法永远无效，也不构成无限盲调的理由。')

page('6 分层接口与训练预算')
p('进入分层阶段前新增高层接口：高层每10个低层步或任务完成时选择公开候选区、已发现目标或继续当前任务；连续施救期间不随意改派。候选列表由已观测历史和未覆盖区域构造，禁止隐藏人员位置。低层仍输出原9动作，环境仍每5秒推进一次。')
p('HATRPO加规划路线先训练任务分配头，低层规划固定。双学习器路线先固定高层训练HASAC目标导航，再固定低层训练高层，最后才做小学习率交替适配。高层任务或动作空间变化时重新初始化相应输出头；不把旧逐步动作权重称为已训练的任务分配器。')
p('训练必须记录实际执行动作及其正确概率、服务/规划/安全接管来源、高层决策区间和区间折扣回报。低层受接管动作不能伪装成神经策略采样；旧经验回放在接口或奖励变化时清空并注明。对未知水情的预测和实际观测分开。')
table(['阶段','预算','进入下一阶段的条件'],[
 ['接口检查','每候选64回合成本及正确性试跑','语义一致，指标可复现，回放和保存恢复正确'],
 ['方案筛选','每候选累计128回合，先1个训练种子','超过匹配纯规划或已有强基线，冲突及获救无退步'],
 ['稳定性确认','有效候选3种训练种子，各累计256回合','三种子分布及差场景支持改善，失败完整保留'],
 ['有条件扩训','各累计512再至1024，每32保存、64验证','连续3次验证无预定收益则停止该候选扩训']],[1.25,2.45,3.3])
p('两条学习路线全部进入时，正式训练上限为2路线×3种子×1024＝6144个任务回合，含已完成的阶段预算。另列最多256回合低层预训练预算，不与联合回合重复加总。筛选种子1001；稳定性种子1001、1002、1003。本轮回放不占训练预算。')

page('7 数据扩展和最终验收')
p('真正扩大训练前，先把独立合成父水场从4个训练扩到16个，从2个开发验证扩到总计6个；若生成成本和内部数值检查允许，再增至32和8个。同一父水场的所有人员布局留在同一分组，参数或空间近重复的来源按组拆分。新场景的有限非负、稳定性、守恒及时间坐标单位检查沿用水情要求。')
p('旧2个开发父水场保留作历史对照，新增开发父场景单列。新旧开发结果不能混成现场准确率。最终50场仅在模型、控制规则、执行方式、奖励、阈值和选择规则全部冻结后运行一次；之后的调参必须另建开发和最终测试划分，不能把50场当调参集。')
table(['要证明的能力','所需对照'],[
 ['规划收益','同一冻结模型，补漏开关配对；纯规划单列'],
 ['学习收益','同一控制条件下训练前后对比；规则接管比例及神经自由动作表现同时报告'],
 ['分层收益','高层规则加低层规划、学习高层加低层规划、双学习器分别比较'],
 ['部署方式','随机执行与固定选动作分开；不能用随机较好成绩代替固定策略检验'],
 ['现实约束','建筑、通信、返航、续航和实际服务动作另登记，逐项有证据后再增加']],[1.25,5.75])
p('结果表必须列明模型父版本、累计与新增回合、联合回合口径、训练种子、独立父水场数、布局数、执行方式、接管比例和失败计数。获救比例不等于全员获救回合比例，全员获救也不等于整场完成。')

page('8 排期交付和依据')
table(['顺序','交付物','成本与完成定义'],[
 ['本轮先执行','冻结计划书与protocol；启动检查；208次完整预选；失败清单与选择结论','回放设30分钟上限；越限保存部分结果并标未完成，不缩分母'],
 ['候选通过后','分层接口、短检查、64及128回合结果','先实测64回合速度，再估后续预算；接口改动时间另计'],
 ['稳定性阶段','扩大场景记录、3训练种子、256至1024阶段表','旧联合试验256回合耗时6.94至10.71分钟，仅作旧代码参考'],
 ['最终冻结','50场最终测试与对照报告','全部失败保留；计划、执行及实测结论分开归档']],[1.1,3.5,2.4])
p('交付至 outputs/uav-plan-preselection-20261101，项目 uav-workbench/docs 同名目录双份归档。完整轨迹、末态及来源哈希保留；新代码隔离，原入口和父检查点可直接回退。')
p('计划冻结后不改门槛。结果另附：无候选通过则停止当前组合扩训，后续分层阶段保留为未完成，不把计划写成已完成训练。')
DOC.add_heading('来源与证据',2)
p('本地依据为20261101的 combinations、joint-training、total256 三份实验报告。用户v0.4协议是讨论稿；采用当前已实现环境并服从六算法禁用要求。')
p('Zhong等，Heterogeneous-Agent Reinforcement Learning，JMLR 2024：https://www.jmlr.org/papers/v25/23-0488.html。顺序更新为设计参考，本项目混合及分层适配不继承其未经核验的理论保证。')
p('Choset与Pignon，Coverage Path Planning，FSR 1997：https://publications.ri.cmu.edu/coverage-path-planning-the-boustrophedon-decomposition。分区与往返覆盖为参考，扇形视野、固定步长、服务和动态水情仍须按本项目验证。')

for root in (DOC.styles.element,DOC._element):
 for border in list(root.xpath('.//w:pBdr')):border.getparent().remove(border)
 for color in root.xpath('.//w:color'):
  color.set(qn('w:val'),'000000')
  for attr in ('themeColor','themeTint','themeShade'):color.attrib.pop(qn('w:'+attr),None)
path=BASE/'HYRA无人机算法预选与训练计划书-v1.0.docx';DOC.save(path)
(BASE/'plan-content.json').write_text(json.dumps(content,ensure_ascii=False,indent=2),encoding='utf8')
print(path)
