import json,math,hashlib,zipfile
from pathlib import Path
import openpyxl
base=Path(__file__).resolve().parent
d=json.loads((base/'data.json').read_text(encoding='utf8'))
p=base/'单双模型训练全量数据-20261101.xlsx'
with zipfile.ZipFile(p) as z:assert z.testzip() is None
w=openpyxl.load_workbook(p,read_only=True,data_only=True)
def equivalent(a,b):
    if b is None:return a is None
    if isinstance(b,(dict,list)):
        try:return json.loads(a)==b
        except (TypeError,ValueError):return False
    if isinstance(b,(int,float)) and not isinstance(b,bool):return isinstance(a,(int,float)) and math.isclose(a,b,rel_tol=1e-12,abs_tol=1e-12)
    return a==b
counts={}
def check(name,expected,start=2):
    rows=list(w[name].iter_rows(min_row=start,max_row=start+len(expected)-1,max_col=len(expected[0]),values_only=True))
    assert len(rows)==len(expected),(name,len(rows),len(expected))
    for i,(got,exp) in enumerate(zip(rows,expected),start):
        for j,(a,b) in enumerate(zip(got,exp),1):assert equivalent(a,b),(name,i,j,str(a)[:140],str(b)[:140])
    counts[name]=len(rows)
ev=[[r['id'],r['experiment'],r['candidate'],r['episode'],r['control'],r['parents'],r['n'],r['completed'],r['completion'],r['RCR'],r['coverage'],r['ACC'],r['ST'],r['max_conflicts'],r['source']] for r in d['evaluations']]
check('全部评估轮次',ev)
fields=['evaluation_id','experiment','candidate','episode','control','case','stratum','layout','assignment','mode','action_seed','row','task_complete','rescued','discovered','RCR','coverage','conflicts','ACC','ST','FRT','CRT','planner_fraction','neural_free_fraction','timed_out','unserved_ids','trace','source']
for current,name in [(True,'当前逐场验证'),(False,'历史逐场验证')]:
    data=[r for r in d['details'] if r['experiment'].startswith(('当前','严格'))==current]
    check(name,[[r.get(k) for k in fields] for r in data])
def train(r):
    return [r.get('experiment'),r.get('candidate'),r.get('joint_episode'),r.get('stage'),r.get('case'),r.get('layout'),r.get('roles'),r.get('optimizer_updates',r.get('updates')),r.get('return_sum'),r.get('task_complete'),r.get('rescued'),r.get('discovered'),r.get('RCR'),r.get('coverage'),r.get('conflicts'),r.get('ACC'),r.get('ST'),r.get('FRT'),r.get('CRT'),r.get('planner_fraction'),r.get('neural_free_fraction'),r.get('elapsed_stage_s',r.get('elapsed_s')),r.get('loss',r.get('diagnostics')),r.get('source'),r.get('source_line')]
for current,name in [(True,'当前逐回合训练'),(False,'历史逐回合训练')]:check(name,[train(r) for r in d['training'] if r['experiment'].startswith('当前')==current])
strict=[[r['candidate'],r['variant'],r['mode'],r['stratum'],r['independent_parents'],r['rollouts'],r['means']['task_complete'],r['means']['RCR'],r['means']['coverage'],r['means']['ACC'],r['means']['ST'],r['failed_rollouts'],r['all_assignments_complete']['completed_parents'],*r['all_assignments_complete']['wilson_ci95'],r['max_conflicts']] for r in d['strict_summary']['summary']]
check('严格测试结果',strict,7)
check('权重检查点',[[r['experiment'],r['candidate'],r['episode'],r['bytes'],r['path']] for r in d['checkpoints']])
evmap={r['id']:r for r in d['evaluations']}
overview=[]
for r in d['overview']:
    best=evmap.get(r['best_id']);latest=evmap.get(r['latest_id'])
    overview.append([r['experiment'].replace('当前',''),r['candidate'].upper().replace('_',' + '),r['actual'],r['target'],r['parent_episodes'],r['best_episode'],best['completion'] if best else None,best['RCR'] if best else None,best['coverage'] if best else None,best['ACC'] if best else None,r['latest_evaluation'],latest['completion'] if latest else None,r['extension']])
check('当前21方案',overview,8)
assert w['当前21方案']['F33'].value==sum(r['actual'] for r in d['overview'])
assert '5630' in w['严格测试结果']['A3'].value
assert d['strict_summary']['rows']==5630 and len(d['details'])==16642 and len(d['training'])==13184
errors=[]
for s in w:
    for row in s.iter_rows():
        for cell in row:
            if cell.data_type=='e':errors.append([s.title,cell.coordinate,cell.value])
assert not errors,errors[:10]
w.close()
result=dict(status='passed',all_exported_table_values_match_source=True,row_counts=counts,formula_caches_verified=True,no_excel_error_cells=True,
    data_sha256=hashlib.sha256((base/'data.json').read_bytes()).hexdigest(),xlsx_sha256=hashlib.sha256(p.read_bytes()).hexdigest(),bytes=p.stat().st_size)
(base/'export-audit.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf8');print(json.dumps(result,ensure_ascii=False))
